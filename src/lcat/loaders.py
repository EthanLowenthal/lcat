"""Turn raw file text into one of a few simple documents the views can render."""

from __future__ import annotations

import csv
import io
import json
import warnings
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import TYPE_CHECKING

if TYPE_CHECKING:  # pragma: no cover
    from PIL.Image import Image as PILImage

from lcat.detect import sniff_delimiter


@dataclass
class MarkdownDoc:
    text: str
    path: Path | None = None


@dataclass
class DelimitedFormat:
    """How a table was parsed out of a CSV/TSV file, so it can be written back."""

    delimiter: str = ","
    has_header: bool = True


@dataclass
class JsonFormat:
    """How a table was flattened out of JSON, so it can be written back."""

    shape: str
    """One of records, rows, scalars, keyvalue."""
    lines: bool = False
    """The source was JSON Lines: one document per line."""
    text_columns: tuple[bool, ...] = ()
    """Per column: the source values were strings, so edits stay strings."""
    text_rows: tuple[bool, ...] = ()
    """Same, per row, for a key/value table: one JSON object has a type per key."""


TableFormat = DelimitedFormat | JsonFormat


@dataclass
class TableDoc:
    columns: list[str]
    rows: list[list[str]] = field(default_factory=list)
    path: Path | None = None
    total_rows: int | None = None
    """Row count before `--max-rows` trimmed it, when it did."""
    fmt: TableFormat = field(default_factory=DelimitedFormat)
    """The source format, used to serialize the table again after an edit."""

    @property
    def shape(self) -> tuple[int, int]:
        return len(self.rows), len(self.columns)

    def head(self, max_rows: int | None) -> "TableDoc":
        """This table, trimmed to `max_rows`, remembering the original count."""
        if max_rows is None or len(self.rows) <= max_rows:
            return self
        return TableDoc(
            self.columns, self.rows[:max_rows], self.path, len(self.rows), self.fmt
        )


@dataclass
class SheetDoc:
    """One worksheet of a workbook, already flattened to strings."""

    name: str
    columns: list[str]
    rows: list[list[str]] = field(default_factory=list)
    total_rows: int | None = None
    """Row count before `--max-rows` trimmed it, when it did."""

    @property
    def shape(self) -> tuple[int, int]:
        return len(self.rows), len(self.columns)

    def head(self, max_rows: int | None) -> "SheetDoc":
        if max_rows is None or len(self.rows) <= max_rows:
            return self
        return SheetDoc(self.name, self.columns, self.rows[:max_rows], len(self.rows))


@dataclass
class WorkbookDoc:
    """A spreadsheet: several sheets, one on screen at a time. Read-only."""

    sheets: list[SheetDoc]
    path: Path | None = None
    index: int = 0
    """Which sheet is on screen."""
    formulas: bool = False
    """Show formulas as written instead of the values Excel last computed."""
    source: bytes | None = None
    """The file itself, kept so the formula view can be parsed on demand."""
    formula_sheets: list[SheetDoc] | None = None
    """The same sheets read as formulas, parsed the first time they are asked for."""
    has_header: bool = True
    max_rows: int | None = None

    @property
    def sheet(self) -> SheetDoc:
        """The sheet on screen, in whichever of the two readings is selected."""
        formulas = self.formulas and self.formula_sheets
        sheets = self.formula_sheets if formulas else self.sheets
        return sheets[min(self.index, len(sheets) - 1)]

    @property
    def names(self) -> list[str]:
        return [sheet.name for sheet in self.sheets]

    def read_formulas(self) -> bool:
        """Parse the formula reading of the workbook. False if it cannot be had."""
        if self.formula_sheets is not None:
            return True
        if self.source is None:
            return False
        self.formula_sheets = _read_sheets(
            self.source, formulas=True, has_header=self.has_header, max_rows=self.max_rows
        )
        return True


@dataclass
class CodeDoc:
    text: str
    lexer: str = "text"
    path: Path | None = None


@dataclass
class ImageDoc:
    """A raster image, shown as pixels where the terminal can and blocks elsewhere."""

    image: PILImage
    path: Path | None = None

    @property
    def format(self) -> str:
        return self.image.format or "image"

    @property
    def size(self) -> tuple[int, int]:
        """(width, height) in pixels."""
        return self.image.width, self.image.height


Document = MarkdownDoc | TableDoc | CodeDoc | ImageDoc | WorkbookDoc


def _cell(value: object) -> str:
    """Render a JSON value as a single-line string."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)):
        return str(value)
    return json.dumps(value, separators=(",", ":"), ensure_ascii=False)


def load_markdown(text: str, path: Path | None = None) -> MarkdownDoc:
    return MarkdownDoc(text, path)


def load_delimited(
    text: str,
    path: Path | None = None,
    delimiter: str | None = None,
    has_header: bool = True,
) -> TableDoc:
    if delimiter is None:
        delimiter = sniff_delimiter(text[:8192])
    fmt = DelimitedFormat(delimiter, has_header)
    raw = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    raw = [row for row in raw if row and any(cell.strip() for cell in row)]
    if not raw:
        return TableDoc([], [], path, fmt=fmt)

    width = max(len(row) for row in raw)
    if has_header:
        header = raw[0] + [""] * (width - len(raw[0]))
        columns = [name.strip() or f"col{i + 1}" for i, name in enumerate(header)]
        body = raw[1:]
    else:
        columns = [f"col{i + 1}" for i in range(width)]
        body = raw

    rows = [row[:width] + [""] * (width - len(row)) for row in body]
    return TableDoc(columns, rows, path, fmt=fmt)


def load_json(text: str, path: Path | None = None) -> TableDoc | CodeDoc:
    data, is_lines = _parse_json(text)
    if data is None:
        return CodeDoc(text, "json", path)

    if isinstance(data, list) and data:
        if all(isinstance(item, dict) for item in data):
            columns: list[str] = []
            seen: set[str] = set()
            for item in data:
                for key in item:
                    if key not in seen:
                        seen.add(key)
                        columns.append(key)
            rows = [[_cell(item.get(key)) for key in columns] for item in data]
            texts = tuple(
                any(isinstance(item.get(key), str) for item in data) for key in columns
            )
            fmt = JsonFormat("records", is_lines, texts)
            return TableDoc(columns, rows, path, fmt=fmt)
        if all(isinstance(item, list) for item in data):
            width = max((len(item) for item in data), default=0)
            columns = [f"col{i + 1}" for i in range(width)]
            rows = [
                [_cell(v) for v in item] + [""] * (width - len(item)) for item in data
            ]
            texts = tuple(
                any(i < len(item) and isinstance(item[i], str) for item in data)
                for i in range(width)
            )
            fmt = JsonFormat("rows", is_lines, texts)
            return TableDoc(columns, rows, path, fmt=fmt)
        if all(not isinstance(item, (dict, list)) for item in data):
            texts = (any(isinstance(item, str) for item in data),)
            fmt = JsonFormat("scalars", is_lines, texts)
            return TableDoc(["value"], [[_cell(item)] for item in data], path, fmt=fmt)

    if isinstance(data, dict) and data and not is_lines:
        rows = [[key, _cell(value)] for key, value in data.items()]
        per_key = tuple(isinstance(value, str) for value in data.values())
        fmt = JsonFormat("keyvalue", False, (True, False), per_key)
        return TableDoc(["key", "value"], rows, path, fmt=fmt)

    pretty = json.dumps(data, indent=2, ensure_ascii=False)
    return CodeDoc(pretty, "json", path)


def _parse_json(text: str) -> tuple[object | None, bool]:
    """Parse JSON or JSON Lines. Returns (data, was_json_lines)."""
    stripped = text.strip()
    if not stripped:
        return None, False
    try:
        return json.loads(stripped), False
    except ValueError:
        pass

    items = []
    for line in stripped.splitlines():
        line = line.strip()
        if not line:
            continue
        try:
            items.append(json.loads(line))
        except ValueError:
            return None, False
    return (items, True) if items else (None, False)


def load_image(data: bytes, path: Path | None = None) -> ImageDoc:
    """Decode an image file. Raises ValueError if PIL cannot read it."""
    from PIL import Image, UnidentifiedImageError

    try:
        image = Image.open(io.BytesIO(data))
        image.load()
    except (UnidentifiedImageError, OSError, ValueError, SyntaxError) as error:
        raise ValueError(f"not a readable image: {error}") from error
    return ImageDoc(image, path)


def _xlsx_cell(value: object) -> str:
    """Render a spreadsheet cell as a single-line string."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        # A date-formatted cell arrives as midnight; show it as a plain date.
        if (value.hour, value.minute, value.second, value.microsecond) == (0, 0, 0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    return str(value)


def _used_width(rows: list[list[str]]) -> int:
    """The last column with anything in it. A sheet usually claims more."""
    width = 0
    for row in rows:
        for index in range(len(row) - 1, width - 1, -1):
            if row[index].strip():
                width = index + 1
                break
    return width


def _sheet(name: str, worksheet, has_header: bool) -> SheetDoc:
    """Flatten one worksheet into a table of strings."""
    from openpyxl.utils import get_column_letter

    rows = [
        [_xlsx_cell(value) for value in row]
        for row in worksheet.iter_rows(values_only=True)
    ]
    # A sheet's stated dimensions usually run past the data: trailing blank rows and
    # columns are dropped so the table is the size it looks in Excel.
    while rows and not any(cell.strip() for cell in rows[-1]):
        rows.pop()
    width = _used_width(rows)
    if not width:
        return SheetDoc(name, [], [])
    rows = [row[:width] + [""] * (width - len(row)) for row in rows]

    letters = [get_column_letter(index + 1) for index in range(width)]
    if has_header:
        columns = [cell.strip() or letters[i] for i, cell in enumerate(rows[0])]
        body = rows[1:]
    else:
        columns, body = letters, rows
    return SheetDoc(name, columns, body)


def _read_sheets(
    data: bytes,
    *,
    formulas: bool = False,
    has_header: bool = True,
    max_rows: int | None = None,
) -> list[SheetDoc]:
    """Read every worksheet of an xlsx. `formulas` reads the text of formula cells
    instead of the values Excel last computed for them."""
    from openpyxl import load_workbook

    try:
        with warnings.catch_warnings():
            # openpyxl warns about parts it drops (data validation, print settings);
            # nothing the reader can act on, and it would garble the display.
            warnings.simplefilter("ignore")
            book = load_workbook(
                io.BytesIO(data), read_only=True, data_only=not formulas
            )
            try:
                return [
                    _sheet(worksheet.title, worksheet, has_header).head(max_rows)
                    for worksheet in book.worksheets
                ]
            finally:
                book.close()
    except (zipfile.BadZipFile, OSError, ValueError, KeyError, TypeError) as error:
        raise ValueError(f"not a readable xlsx: {error}") from error


def load_xlsx(
    data: bytes,
    path: Path | None = None,
    *,
    has_header: bool = True,
    max_rows: int | None = None,
) -> WorkbookDoc:
    """Read a workbook from file bytes. Raises ValueError if openpyxl cannot."""
    sheets = _read_sheets(data, has_header=has_header, max_rows=max_rows)
    if not sheets:
        raise ValueError("the workbook has no sheets")
    return WorkbookDoc(
        sheets, path, source=data, has_header=has_header, max_rows=max_rows
    )


def load(
    text: str,
    mode: str,
    path: Path | None = None,
    delimiter: str | None = None,
    has_header: bool = True,
) -> Document:
    if mode == "md":
        return load_markdown(text, path)
    if mode == "json":
        return load_json(text, path)
    if mode == "tsv":
        return load_delimited(text, path, delimiter or "\t", has_header)
    if mode == "csv":
        return load_delimited(text, path, delimiter, has_header)
    if mode == "img":
        raise ValueError("images are loaded from bytes with load_image()")
    if mode == "xlsx":
        raise ValueError("workbooks are loaded from bytes with load_xlsx()")
    raise ValueError(f"unknown mode: {mode!r}")
