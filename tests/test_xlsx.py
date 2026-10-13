import contextlib
import io
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook
from rich.console import Console
from textual.coordinate import Coordinate
from textual.widgets import DataTable

from lcat.app import LcatApp
from lcat.detect import is_xlsx, mode_from_extension, resolve_mode
from lcat.loaders import WorkbookDoc, load, load_xlsx
from lcat.plain import render
from lcat.save import SaveError, serialize
from lcat.views import WorkbookView


def workbook_bytes(sheets: dict[str, list[list]]) -> bytes:
    """An xlsx holding `sheets`, as file bytes."""
    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


SHEETS = {
    "Sales": [["region", "units"], ["North", 120], ["South", 84]],
    "Notes": [["field", "value"], ["owner", "planning"]],
}


def sample() -> bytes:
    return workbook_bytes(SHEETS)


# -- detection -------------------------------------------------------------


def test_extension_mapping():
    assert mode_from_extension(Path("book.xlsx")) == "xlsx"
    assert mode_from_extension(Path("book.XLSM")) == "xlsx"


def test_is_xlsx_recognises_a_workbook_and_not_other_zips():
    assert is_xlsx(sample())
    assert not is_xlsx(b"PK\x03\x04" + b"just a zip of something else")
    assert not is_xlsx(b"region,units\nNorth,120\n")


def test_resolve_mode_sniffs_a_workbook_without_an_extension():
    raw = sample()[:8192]
    assert resolve_mode(Path("book"), "auto", raw.decode("utf-8", "replace"), raw) == "xlsx"
    assert resolve_mode(None, "auto", raw.decode("utf-8", "replace"), raw) == "xlsx"


# -- loading ---------------------------------------------------------------


def test_sheets_are_loaded_in_order_with_headers():
    doc = load_xlsx(sample())
    assert isinstance(doc, WorkbookDoc)
    assert doc.names == ["Sales", "Notes"]
    assert doc.sheet.name == "Sales"
    assert doc.sheet.columns == ["region", "units"]
    assert doc.sheet.rows == [["North", "120"], ["South", "84"]]
    assert doc.sheet.shape == (2, 2)


def test_values_are_flattened_to_single_line_strings():
    doc = load_xlsx(
        workbook_bytes(
            {
                "S": [
                    ["when", "at", "flag", "empty", "num"],
                    [date(2026, 3, 4), datetime(2026, 3, 4, 9, 30), True, None, 22.0],
                ]
            }
        )
    )
    assert doc.sheet.rows == [["2026-03-04", "2026-03-04 09:30:00", "TRUE", "", "22"]]


def test_trailing_blank_rows_and_columns_are_dropped():
    doc = load_xlsx(
        workbook_bytes({"S": [["a", "b", None, None], ["1", "2"], [None, None]]})
    )
    assert doc.sheet.columns == ["a", "b"]
    assert doc.sheet.rows == [["1", "2"]]


def test_gaps_inside_the_sheet_are_kept_and_unnamed_columns_get_letters():
    doc = load_xlsx(workbook_bytes({"S": [["a", None, "c"], ["1", None, "3"]]}))
    assert doc.sheet.columns == ["a", "B", "c"]
    assert doc.sheet.rows == [["1", "", "3"]]


def test_no_header_names_columns_by_spreadsheet_letter():
    doc = load_xlsx(sample(), has_header=False)
    assert doc.sheet.columns == ["A", "B"]
    assert doc.sheet.rows[0] == ["region", "units"]


def test_max_rows_trims_every_sheet_and_remembers_the_count():
    doc = load_xlsx(sample(), max_rows=1)
    assert doc.sheet.rows == [["North", "120"]]
    assert doc.sheet.total_rows == 2


def test_an_empty_sheet_loads_as_an_empty_table():
    doc = load_xlsx(workbook_bytes({"Blank": []}))
    assert doc.sheets[0].columns == []
    assert doc.sheets[0].rows == []


def test_formulas_are_read_on_demand():
    doc = load_xlsx(workbook_bytes({"S": [["a", "double"], [2, "=A2*2"]]}))
    # Nothing has computed the formula, so the value reading of the cell is empty.
    assert doc.sheet.rows == [["2", ""]]
    assert doc.read_formulas()
    doc.formulas = True
    assert doc.sheet.rows == [["2", "=A2*2"]]


def test_formulas_need_the_source_bytes():
    doc = load_xlsx(sample())
    doc.source = None
    assert not doc.read_formulas()


def test_a_file_that_is_not_a_workbook_is_rejected():
    with pytest.raises(ValueError, match="not a readable xlsx"):
        load_xlsx(b"PK\x03\x04 not really a workbook")


def test_load_points_at_the_bytes_loader():
    with pytest.raises(ValueError, match="load_xlsx"):
        load("", "xlsx")


# -- plain and save --------------------------------------------------------


def test_plain_render_prints_every_sheet():
    buffer = io.StringIO()
    with contextlib.redirect_stdout(buffer):
        render(load_xlsx(sample()))
    out = buffer.getvalue()
    assert "Sales" in out and "North" in out
    assert "Notes" in out and "planning" in out


def test_workbooks_cannot_be_written_back():
    with pytest.raises(SaveError, match="read-only"):
        serialize(load_xlsx(sample()))


# -- the interactive view --------------------------------------------------


def notes(app: LcatApp) -> list[str]:
    seen: list[str] = []
    app.notify = lambda message, **kwargs: seen.append(message)
    return seen


async def test_the_first_sheet_is_shown_with_its_name():
    app = LcatApp(load_xlsx(sample()))
    async with app.run_test():
        table = app.query_one(DataTable)
        assert [column.label.plain for column in table.ordered_columns] == [
            "region",
            "units",
        ]
        assert table.row_count == 2
        assert "Sales" in app.sub_title


async def test_brackets_walk_the_sheets():
    app = LcatApp(load_xlsx(sample()))
    async with app.run_test() as pilot:
        view = app.query_one(WorkbookView)
        await pilot.press("right_square_bracket")
        await pilot.pause()
        assert view.workbook.index == 1
        assert "Notes" in app.sub_title
        assert app.query_one(DataTable).row_count == 1
        await pilot.press("left_square_bracket")
        await pilot.pause()
        assert view.workbook.index == 0
        assert "Sales" in app.sub_title


async def test_switching_sheets_puts_the_cursor_back_at_the_top():
    app = LcatApp(load_xlsx(sample()))
    async with app.run_test() as pilot:
        await pilot.press("G", "dollar_sign")
        assert app.query_one(DataTable).cursor_coordinate != Coordinate(0, 0)
        await pilot.press("right_square_bracket")
        await pilot.pause()
        assert app.query_one(DataTable).cursor_coordinate == Coordinate(0, 0)


async def test_one_sheet_workbook_says_so_instead_of_wrapping():
    app = LcatApp(load_xlsx(workbook_bytes({"Only": [["a"], ["1"]]})))
    async with app.run_test() as pilot:
        seen = notes(app)
        await pilot.press("right_square_bracket")
        await pilot.pause()
        assert seen == ["the workbook has one sheet"]


async def test_f_toggles_formulas_and_keeps_the_cursor():
    app = LcatApp(load_xlsx(workbook_bytes({"S": [["a", "double"], [2, "=A2*2"]]})))
    async with app.run_test() as pilot:
        view = app.query_one(WorkbookView)
        await pilot.press("right")
        assert view.cell_value(Coordinate(0, 1)) == ""
        await pilot.press("f")
        await pilot.pause()
        assert view.cell_value(Coordinate(0, 1)) == "=A2*2"
        assert app.query_one(DataTable).cursor_coordinate == Coordinate(0, 1)
        assert "formulas" in app.sub_title
        await pilot.press("f")
        await pilot.pause()
        assert view.cell_value(Coordinate(0, 1)) == ""
        assert "formulas" not in app.sub_title


async def test_editing_and_saving_are_refused():
    app = LcatApp(load_xlsx(sample()))
    async with app.run_test() as pilot:
        seen = notes(app)
        await pilot.press("i")
        await pilot.press("ctrl+s")
        await pilot.pause()
        assert seen == ["xlsx files are read-only"] * 2
        assert not app.dirty


async def test_reload_keeps_the_sheet_and_the_formula_view(tmp_path):
    path = tmp_path / "book.xlsx"
    path.write_bytes(workbook_bytes({"One": [["a"], [1]], "Two": [["b"], ["=1+1"]]}))

    def reload():
        return load_xlsx(path.read_bytes(), path)

    app = LcatApp(reload(), reload=reload)
    async with app.run_test() as pilot:
        await pilot.press("right_square_bracket", "f")
        await pilot.pause()
        assert app.doc.sheet.rows == [["=1+1"]]
        path.write_bytes(workbook_bytes({"One": [["a"], [1]], "Two": [["b"], ["=2+2"]]}))
        await app._poll_disk()
        await pilot.pause()
        assert app.doc.names[app.doc.index] == "Two"
        assert app.doc.formulas
        assert app.doc.sheet.rows == [["=2+2"]]


async def test_reload_falls_back_to_the_first_sheet_when_one_disappears(tmp_path):
    path = tmp_path / "book.xlsx"
    path.write_bytes(sample())

    def reload():
        return load_xlsx(path.read_bytes(), path)

    app = LcatApp(reload(), reload=reload)
    async with app.run_test() as pilot:
        await pilot.press("right_square_bracket")
        await pilot.pause()
        path.write_bytes(workbook_bytes({"Sales": SHEETS["Sales"]}))
        await app._poll_disk()
        await pilot.pause()
        assert app.doc.index == 0
        assert "Sales" in app.sub_title


# -- formatting --------------------------------------------------------------


def styled_bytes(
    rows, *, fonts=None, fills=None, formats=None, aligns=None, widths=None
) -> bytes:
    """An xlsx whose cells carry formatting, keyed by coordinate: fonts={"A1": Font(...)}."""
    from openpyxl.styles import Alignment

    book = Workbook()
    sheet = book.active
    sheet.title = "S"
    for row in rows:
        sheet.append(row)
    for coord, font in (fonts or {}).items():
        sheet[coord].font = font
    for coord, fill in (fills or {}).items():
        sheet[coord].fill = fill
    for coord, fmt in (formats or {}).items():
        sheet[coord].number_format = fmt
    for coord, how in (aligns or {}).items():
        sheet[coord].alignment = Alignment(horizontal=how)
    for letter, width in (widths or {}).items():
        sheet.column_dimensions[letter].width = width
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def blue_header() -> bytes:
    from openpyxl.styles import Font, PatternFill

    return styled_bytes(
        [["region", "revenue"], ["North", 1234.5], ["South", -800]],
        fonts={"A1": Font(bold=True, color="FFFFFFFF"), "B1": Font(bold=True)},
        fills={
            "A1": PatternFill("solid", fgColor="FF4F81BD"),
            "A3": PatternFill("solid", fgColor="FFFFFF00"),
        },
        formats={"B2": '"$"#,##0.00', "B3": '"$"#,##0.00'},
        widths={"A": 20},
    )


def test_styles_line_up_with_the_rows_they_describe():
    sheet = load_xlsx(blue_header()).sheet
    assert len(sheet.styles) == len(sheet.rows)
    assert all(len(s) == len(r) for s, r in zip(sheet.styles, sheet.rows))
    assert len(sheet.column_styles) == len(sheet.columns)
    assert sheet.colored


def test_the_header_carries_its_fill_and_font():
    sheet = load_xlsx(blue_header()).sheet
    header = sheet.column_styles[0]
    assert header.background == "#4f81bd"
    assert header.color == "#ffffff"
    assert header.bold


def test_number_formats_reach_the_cell_text():
    sheet = load_xlsx(blue_header()).sheet
    assert sheet.rows[0][1] == "$1,234.50"
    assert sheet.rows[1][1] == "-$800.00"


def test_numbers_are_right_aligned_and_text_is_not():
    sheet = load_xlsx(blue_header()).sheet
    assert sheet.styles[0][0].align == "left"
    assert sheet.styles[0][1].align == "right"


def test_excel_column_widths_are_carried():
    sheet = load_xlsx(blue_header()).sheet
    assert sheet.widths[0] == 20
    assert sheet.widths[1] is None


def test_cells_sharing_a_format_share_one_style_object():
    """The intern cache is what keeps a big sheet from paying per cell."""
    rows = [["a", "b"]] + [["x", "y"] for _ in range(20)]
    sheet = load_xlsx(styled_bytes(rows)).sheet
    first = sheet.styles[0][0]
    assert all(style is first for row in sheet.styles for style in row)


def test_a_plain_workbook_has_no_colours():
    """Alignment alone must not move a sheet onto the white canvas."""
    sheet = load_xlsx(sample()).sheet
    assert not sheet.colored
    assert all(
        style.color is None and style.background is None
        for row in sheet.styles
        for style in row
    )


def test_an_unfilled_cell_keeps_no_background():
    sheet = load_xlsx(blue_header()).sheet
    assert sheet.styles[0][0].background is None


def test_head_slices_the_styles_with_the_rows():
    doc = load_xlsx(blue_header(), max_rows=1)
    sheet = doc.sheet
    assert len(sheet.rows) == 1 and len(sheet.styles) == 1
    assert sheet.column_styles is not None and sheet.widths is not None


def test_no_style_drops_the_formatting_but_not_the_values():
    doc = load_xlsx(blue_header(), styles=False)
    sheet = doc.sheet
    assert sheet.styles is None and sheet.widths is None and not sheet.colored
    assert sheet.rows[0][1] == "1234.5"  # unformatted, as before
    assert doc.read_formulas()
    assert doc.formula_sheets[0].styles is None


# -- formatting on screen ----------------------------------------------------


async def test_cells_reach_the_table_with_their_styling():
    app = LcatApp(load_xlsx(blue_header()))
    async with app.run_test():
        table = app.query_one(DataTable)
        cell = table.get_cell_at(Coordinate(0, 1))
        assert cell.text == "$1,234.50"
        assert cell.style.align == "right"
        assert table.ordered_columns[0].width == 20


async def test_a_coloured_sheet_is_drawn_on_paper():
    app = LcatApp(load_xlsx(blue_header()))
    async with app.run_test(size=(40, 10)):
        view = app.query_one("#view")
        assert view.has_class("styled")
        table = app.query_one(DataTable)
        backgrounds = set()
        for y in range(3):
            for segment in table.render_line(y):
                color = segment.style.bgcolor if segment.style else None
                if color is not None and color.triplet is not None:
                    backgrounds.add(color.triplet.hex)
        assert "#ffffff" in backgrounds  # the paper, untinted
        assert "#4f81bd" in backgrounds  # the header's own fill


async def test_a_plain_sheet_keeps_the_app_theme():
    app = LcatApp(load_xlsx(sample()))
    async with app.run_test():
        assert not app.query_one("#view").has_class("styled")


async def test_no_style_hands_the_table_plain_strings():
    app = LcatApp(load_xlsx(blue_header(), styles=False))
    async with app.run_test():
        table = app.query_one(DataTable)
        assert table.get_cell_at(Coordinate(0, 1)) == "1234.5"
        assert not app.query_one("#view").has_class("styled")


async def test_sorting_a_styled_sheet_keeps_working():
    """DataTable.sort hands the stored cell to the key function, renderable and all."""
    app = LcatApp(load_xlsx(blue_header()))
    async with app.run_test() as pilot:
        await pilot.press("right", "s")
        await pilot.pause()
        table = app.query_one(DataTable)
        # -$800.00 sorts below $1,234.50 despite the currency and the sign.
        assert str(table.get_cell_at(Coordinate(0, 1))) == "-$800.00"


def test_styling_reaches_a_pipe():
    console = Console(
        width=50, force_terminal=True, color_system="truecolor", highlight=False
    )
    with console.capture() as capture:
        render(load_xlsx(blue_header()), console=console)
    out = capture.get()
    assert "48;2;79;129;189" in out  # the header fill
    assert "48;2;255;255;255" in out  # the paper
    assert "$1,234.50" in out
